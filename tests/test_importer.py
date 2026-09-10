import io
import unittest
from datetime import datetime
import openpyxl
import xlwt
from importer import Table, read_table, normalize, export_csv, suggest_mapping


class ImportTests(unittest.TestCase):
    def test_french_csv_preserves_identifier_and_reports_every_invalid_row(self):
        table = read_table('Référence;Prix;Date\n0012;1 234,50;10/09/2026\n0013;12,34,56;31/02/2026\n'.encode(), 'x.csv')
        schema = [{'name': 'id', 'aliases': ['Reference'], 'required': True},
                  {'name': 'amount', 'aliases': ['Prix'], 'type': 'decimal', 'decimal': ',', 'group': ' '},
                  {'name': 'date', 'type': 'date', 'formats': ['%d/%m/%Y']}]
        result = normalize(table, schema)
        self.assertEqual(result['total'], 2)
        self.assertEqual(result['invalid'], 1)
        self.assertEqual(result['rows'][0]['normalized'], {'id': '0012', 'amount': '1234.50', 'date': '2026-09-10'})
        self.assertEqual(result['rows'][1]['raw']['Prix'], '12,34,56')
        self.assertEqual(len(result['rows'][1]['errors']), 2)

    def test_ambiguous_dates_not_guessed(self):
        result = normalize(Table(['d'], [['01/02/2026'], ['13/02/2026']], []),
                           [{'name': 'd', 'type': 'date', 'formats': ['%d/%m/%Y', '%m/%d/%Y']}])
        self.assertEqual(result['rows'][0]['errors'][0]['message'], 'Ambiguous date')
        self.assertEqual(result['rows'][1]['normalized']['d'], '2026-02-13')

    def test_real_xlsx_dates_and_formulas(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(['id', 'date', 'amount'])
        sheet.append(['0007', datetime(2026, 9, 10), '=1+2'])
        data = io.BytesIO(); book.save(data); book.close()
        result = normalize(read_table(data.getvalue(), 'file.xlsx'),
                           [{'name': 'id'}, {'name': 'date', 'type': 'date'}, {'name': 'amount', 'type': 'decimal'}])
        self.assertEqual(result['rows'][0]['normalized']['date'], '2026-09-10')
        self.assertEqual(result['rows'][0]['raw']['amount'], {'formula': '=1+2'})
        self.assertEqual(result['invalid'], 1)

    def test_real_legacy_xls(self):
        book = xlwt.Workbook(); sheet = book.add_sheet('Input')
        for i, h in enumerate(['id', 'date', 'amount']): sheet.write(0, i, h)
        sheet.write(1, 0, '0008')
        sheet.write(1, 1, datetime(2026, 9, 10), xlwt.easyxf(num_format_str='YYYY-MM-DD'))
        sheet.write(1, 2, 12.5)
        data = io.BytesIO(); book.save(data)
        result = normalize(read_table(data.getvalue(), 'file.xls'),
                           [{'name': 'id'}, {'name': 'date', 'type': 'date'}, {'name': 'amount', 'type': 'decimal'}])
        self.assertEqual(result['rows'][0]['normalized'], {'id': '0008', 'date': '2026-09-10', 'amount': '12.5'})
        self.assertIn('cached', result['notes'][0])

    def test_duplicate_headers_and_extra_cells_are_rejected(self):
        for content in ['id;id\na;b\n', 'a;b\n1;2;3\n']:
            with self.assertRaises(ValueError): read_table(content.encode(), 'x.csv', delimiter=';')

    def test_missing_fields_and_blank_rows_not_discarded(self):
        result = normalize(read_table(b'a;b\n;\n1;2\n', 'x.csv'), [{'name': 'required', 'required': True}])
        self.assertEqual(result['total'], 2)
        self.assertEqual(result['invalid'], 2)

    def test_mapping_ambiguity_and_manual_override(self):
        table = Table(['Nom', 'nom'], [['first', 'second']], [])
        rules = [{'name': 'name', 'aliases': ['nom']}]
        self.assertIsNone(suggest_mapping(table.headers, rules)['name'])
        self.assertEqual(normalize(table, rules, {'name': 'nom'})['rows'][0]['normalized']['name'], 'second')
        with self.assertRaises(ValueError): normalize(table, rules, {'name': 'missing'})

    def test_number_grouping_and_finite_values(self):
        table = Table(['amount'], [['1,234.50'], ['12,34.50'], [float('inf')], ['NaN']], [])
        result = normalize(table, [{'name': 'amount', 'type': 'decimal', 'group': ','}])
        self.assertEqual(result['invalid'], 3)
        self.assertEqual(result['rows'][0]['normalized']['amount'], '1234.50')

    def test_duration_and_status(self):
        result = normalize(Table(['duration', 'state'], [['25:30', 'OK'], ['1:99', 'unknown']], []),
                           [{'name': 'duration', 'type': 'duration', 'unit': 'hh:mm'},
                            {'name': 'state', 'type': 'status', 'values': {'OK': 'approved'}}])
        self.assertEqual(result['rows'][0]['normalized']['duration'], '91800')
        self.assertEqual(len(result['rows'][1]['errors']), 2)

    def test_csv_export_neutralizes_formula_text(self):
        result = normalize(Table(['text'], [['=HYPERLINK("https://example.com")'], ['  @SUM(1)']], []), [{'name': 'text'}])
        text = export_csv(result)
        self.assertIn("'=HYPERLINK", text)
        self.assertIn("'@SUM", text)
        self.assertEqual(len(text.splitlines()), 3)

    def test_encoding_is_explicit_and_strict(self):
        data = 'id;label\n001;café\n'.encode('cp1252')
        with self.assertRaises(UnicodeDecodeError): read_table(data, 'x.csv')
        self.assertEqual(read_table(data, 'x.csv', encoding='cp1252').rows[0][1], 'café')

    def test_multiple_sheets_require_selection(self):
        book = openpyxl.Workbook(); book.active.append(['id']); book.create_sheet('Two').append(['id'])
        data = io.BytesIO(); book.save(data); book.close()
        with self.assertRaisesRegex(ValueError, 'Select sheet'): read_table(data.getvalue(), 'x.xlsx')
        self.assertEqual(read_table(data.getvalue(), 'x.xlsx', sheet='Two').headers, ['id'])


if __name__ == '__main__':
    unittest.main()
