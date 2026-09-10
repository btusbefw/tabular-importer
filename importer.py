"""Deterministic tabular ingestion. No network, persistence or formula evaluation."""
import argparse
import csv
import io
import json
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLS = 200


@dataclass
class Table:
    headers: list[str]
    rows: list[list]
    notes: list[str]


def read_table(data: bytes, filename: str, *, sheet=None, encoding="utf-8-sig", delimiter=None):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("File must contain 1 byte to 10 MiB")
    ext = Path(filename).suffix.lower()
    notes = []
    if ext == ".csv":
        text = data.decode(encoding, errors="strict")
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
            except csv.Error:
                raise ValueError("Select the CSV delimiter explicitly") from None
        if delimiter not in [",", ";", "\t", "|"]:
            raise ValueError("Unsupported CSV delimiter")
        source = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        rows = []
        for row in source:
            rows.append(row)
            if len(rows) > MAX_ROWS + 1 or len(row) > MAX_COLS:
                raise ValueError("Limit: 10,000 data rows and 200 columns")
    elif ext == ".xlsx":
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist()) > 2000 or sum(f.file_size for f in z.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds limits")
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ValueError("Select sheet: " + ", ".join(book.sheetnames))
            ws = book[sheet or book.sheetnames[0]]
            # Ignore advertised dimensions; enforce limits while actually reading cells.
            ws.reset_dimensions()
            rows = []
            for row in ws.iter_rows():
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        values.append({"formula": cell.value})
                    elif cell.data_type == "e":
                        values.append({"excel_error": cell.value})
                    else:
                        values.append(cell.value)
                rows.append(values)
                if len(rows) > MAX_ROWS + 1 or len(values) > MAX_COLS:
                    raise ValueError("Limit: 10,000 data rows and 200 columns")
        finally:
            book.close()
        notes.append("XLSX formulas are preserved and flagged; never evaluated.")
    elif ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data, on_demand=True, logfile=io.StringIO())
        try:
            if sheet is None and book.nsheets != 1:
                raise ValueError("Select sheet: " + ", ".join(book.sheet_names()))
            ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
            if ws.nrows > MAX_ROWS + 1 or ws.ncols > MAX_COLS:
                raise ValueError("Limit: 10,000 data rows and 200 columns")
            rows = []
            for r in range(ws.nrows):
                row = []
                for cell in ws.row(r):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                        if 0 <= cell.value < 1:
                            value = value.time()
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = {"excel_error": xlrd.error_text_from_code.get(cell.value, "Unknown")}
                    else:
                        value = cell.value
                    row.append(value)
                rows.append(row)
        finally:
            book.release_resources()
        notes.append("Legacy XLS exposes cached values; formula provenance cannot be established by this reader. No formulas are executed.")
    else:
        raise ValueError("Supported extensions: .csv, .xlsx, .xls")
    if not rows:
        raise ValueError("Missing header row")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    if not headers or any(not v for v in headers) or len(set(headers)) != len(headers):
        raise ValueError("Headers must be nonempty and unique")
    body = []
    for row in rows[1:]:
        if len(row) > len(headers):
            raise ValueError("Row has more cells than the header; fix the source instead of dropping data")
        body.append(row + [None] * (len(headers) - len(row)))
    return Table(headers, body, notes)


def canonical(text):
    return "".join(c for c in unicodedata.normalize("NFKD", text).casefold() if c.isalnum())


def validate_schema(schema):
    if not isinstance(schema, list) or not 1 <= len(schema) <= MAX_COLS:
        raise ValueError("Schema must contain 1–200 fields")
    names = []
    for rule in schema:
        if not isinstance(rule, dict) or not isinstance(rule.get("name"), str) or not rule["name"].strip():
            raise ValueError("Each rule needs a nonempty name")
        names.append(rule["name"])
        if len(rule["name"]) > 256 or not isinstance(rule.get("required", False), bool):
            raise ValueError("Invalid field name or required flag")
        if rule.get("type", "text") not in {"text", "decimal", "date", "duration", "status"}:
            raise ValueError("Unsupported field type")
        if not isinstance(rule.get("aliases", []), list) or any(not isinstance(a, str) for a in rule.get("aliases", [])):
            raise ValueError("Aliases must be text")
        formats = rule.get("formats", ["%Y-%m-%d"])
        if not isinstance(formats, list) or not formats or any(not isinstance(f, str) or len(f) > 100 for f in formats):
            raise ValueError("Date formats must be a nonempty list of strings")
        if any(re.search(r"%[HIMSpfzZXc]", f) for f in formats):
            raise ValueError("Date rules cannot silently discard time or timezone components")
        values = rule.get("values", {})
        if not isinstance(values, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in values.items()):
            raise ValueError("Status values must map text to text")
        for bound in ["min", "max"]:
            if bound in rule:
                try:
                    valid_bound = not isinstance(rule[bound], bool) and Decimal(str(rule[bound])).is_finite()
                except InvalidOperation:
                    valid_bound = False
                if not valid_bound:
                    raise ValueError("Numeric bounds must be finite decimals")
        if "min" in rule and "max" in rule and Decimal(str(rule["min"])) > Decimal(str(rule["max"])):
            raise ValueError("Minimum exceeds maximum")
        if rule.get("decimal", ".") not in {".", ","}:
            raise ValueError("Decimal separator must be . or ,")
        group = rule.get("group", "")
        if group not in {"", " ", ".", ","} or group == rule.get("decimal", "."):
            raise ValueError("Invalid grouping separator")
        if rule.get("type") == "duration" and rule.get("unit", "minutes") not in {"minutes", "hours", "seconds", "hh:mm"}:
            raise ValueError("Invalid duration unit")
    if len(set(names)) != len(names):
        raise ValueError("Duplicate target field names")


def suggest_mapping(headers, schema):
    validate_schema(schema)
    mapping = {}
    for rule in schema:
        labels = {canonical(s) for s in [rule["name"], *rule.get("aliases", [])]}
        matches = [h for h in headers if canonical(h) in labels]
        mapping[rule["name"]] = matches[0] if len(matches) == 1 else None
    return mapping


def decimal_value(value, rule):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        result = Decimal(str(value))
    else:
        text = str(value).strip().replace("\u00a0", " ").replace("\u202f", " ")
        dec, group = rule.get("decimal", "."), rule.get("group", "")
        integer = r"\d+" if not group else rf"(?:\d+|\d{{1,3}}(?:{re.escape(group)}\d{{3}})+)"
        if not re.fullmatch(rf"[+-]?{integer}(?:{re.escape(dec)}\d+)?", text):
            raise ValueError("Invalid number for the selected separators")
        result = Decimal(text.replace(group, "") .replace(dec, ".") if group else text.replace(dec, "."))
    if not result.is_finite():
        raise ValueError("Non-finite number")
    if "min" in rule and result < Decimal(str(rule["min"])):
        raise ValueError("Below minimum")
    if "max" in rule and result > Decimal(str(rule["max"])):
        raise ValueError("Above maximum")
    return result


def convert(value, rule):
    if isinstance(value, dict):
        raise ValueError("Spreadsheet formula/error requires review")
    if value is None or isinstance(value, str) and not value.strip():
        if rule.get("required", False):
            raise ValueError("Required value missing")
        return None
    kind = rule.get("type", "text")
    if kind == "text":
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()
    if kind == "decimal":
        return format(decimal_value(value, rule), "f")
    if kind == "date":
        if isinstance(value, (date, datetime)):
            if isinstance(value, datetime) and value.time().isoformat() != "00:00:00":
                raise ValueError("Date contains time; explicit rule required before truncating")
            return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
        matches = set()
        for pattern in rule.get("formats", ["%Y-%m-%d"]):
            try:
                matches.add(datetime.strptime(str(value).strip(), pattern).date().isoformat())
            except ValueError:
                pass
        if len(matches) != 1:
            raise ValueError("Ambiguous date" if matches else "Invalid date for selected formats")
        return matches.pop()
    if kind == "status":
        text = str(value).strip()
        values = rule.get("values", {})
        if text not in values:
            raise ValueError("Unknown status")
        return values[text]
    if kind == "duration":
        unit = rule.get("unit", "minutes")
        if isinstance(value, time):
            if value.tzinfo:
                raise ValueError("Duration cannot contain a timezone")
            result = Decimal(value.hour * 3600 + value.minute * 60 + value.second) + Decimal(value.microsecond) / 1000000
        elif isinstance(value, timedelta):
            result = Decimal(value.days * 86400 + value.seconds) + Decimal(value.microseconds) / 1000000
        elif unit == "hh:mm":
            match = re.fullmatch(r"(\d+):([0-5]\d)", str(value).strip())
            if not match:
                raise ValueError("Expected hours:minutes, e.g. 02:30")
            result = Decimal(int(match[1]) * 3600 + int(match[2]) * 60)
        else:
            result = decimal_value(value, rule) * {"hours": 3600, "minutes": 60, "seconds": 1}[unit]
        if result < 0:
            raise ValueError("Negative duration")
        return format(result, "f")


def normalize(table, schema, mapping=None):
    validate_schema(schema)
    mapping = suggest_mapping(table.headers, schema) if mapping is None else mapping
    names = {rule["name"] for rule in schema}
    if not isinstance(mapping, dict) or set(mapping) != names or any(h is not None and (not isinstance(h, str) or h not in table.headers) for h in mapping.values()):
        raise ValueError("Mapping must specify every target using a source header or null")
    results = []
    for number, row in enumerate(table.rows, 2):
        def serializable(value):
            if isinstance(value, (datetime, date, time)):
                return value.isoformat()
            if isinstance(value, timedelta):
                return {"duration_seconds": str(value.total_seconds())}
            if isinstance(value, float) and not math.isfinite(value):
                return {"non_finite": str(value)}
            return value
        raw = {h: serializable(v) for h, v in zip(table.headers, row)}
        record, errors = {}, []
        for rule in schema:
            target = rule["name"]
            source = mapping[target]
            value = row[table.headers.index(source)] if source else None
            try:
                record[target] = convert(value, rule)
            except (ValueError, InvalidOperation) as error:
                record[target] = None
                errors.append({"field": target, "source": source, "message": str(error)})
        results.append({"row": number, "raw": raw, "normalized": record, "errors": errors})
    return {"mapping": mapping, "notes": table.notes, "total": len(results),
            "invalid": sum(bool(r["errors"]) for r in results), "rows": results}


def export_csv(result):
    """All rows, including invalid rows. Spreadsheet-safe escaping is explicit."""
    fields = list(result["mapping"])
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    def safe(value):
        text = "" if value is None else str(value)
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")) else text
    writer.writerow(["_source_row", "_errors", *map(safe, fields)])
    for row in result["rows"]:
        writer.writerow([row["row"], safe(json.dumps(row["errors"], ensure_ascii=False)),
                         *[safe(row["normalized"][f]) for f in fields]])
    return out.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("schema", type=Path)
    parser.add_argument("--mapping", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--encoding", default="utf-8-sig")
    parser.add_argument("--delimiter")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    with args.input.open("rb") as f:
        data = f.read(MAX_BYTES + 1)
    table = read_table(data, args.input.name, sheet=args.sheet, encoding=args.encoding, delimiter=args.delimiter)
    result = normalize(table, json.loads(args.schema.read_text()),
                       json.loads(args.mapping.read_text()) if args.mapping else None)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    print(f"{result['total']} rows; {result['invalid']} require review. Output: {args.output}")


if __name__ == "__main__":
    main()
